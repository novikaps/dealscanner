"""
Builds the 2-3 page first-screen memo as a Word file (python-docx, offline).
Layout: header banner -> snapshot -> 8 sections -> gaps & caveats.
"""
import datetime as dt
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
HEAD_FILL = "1F3A5F"
BAND_FILL = "EEF2F7"
VERDICT_FILL = {"Proceed to next stage": "E3F1E6", "Proceed with caution": "FFF4DC", "Pass": "FBE3E3"}
NUMERIC = re.compile(r"^(-?[\d,.]+ ?(%|x)?|–)$")  # right-align numbers only
SEVERITY_FILL = {"High": "FBE3E3", "Medium": "FFF4DC", "Low": "E3F1E6"}


# ---------- small helpers ----------

def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")   # 'clear' (never 'solid', which renders black)
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_text(cell, text, bold=False, size=8.5, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    run = p.add_run(str(text) if text not in (None, "") else "–")
    run.bold, run.font.size = bold, Pt(size)
    if color:
        run.font.color.rgb = color


def _table(doc, header, rows, widths_cm, fills=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(c, HEAD_FILL)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            right = i > 0 and isinstance(v, str) and bool(NUMERIC.match(v))
            _cell_text(cells[i], v, align=WD_ALIGN_PARAGRAPH.RIGHT if right else None)
            if fills and fills[r_i] and fills[r_i][i]:
                _shade(cells[i], fills[r_i][i])
    _set_widths(t, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def _set_widths(table, widths_cm):
    """Widths must be set on the grid AND on every cell, or Word/LibreOffice ignore them."""
    for i, w in enumerate(widths_cm):
        table.columns[i].width = Cm(w)
        for cell in table.columns[i].cells:
            cell.width = Cm(w)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(3)
    p.paragraph_format.keep_with_next = True   # never strand a heading at a page bottom
    r = p.add_run(text.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), NAVY
    # bottom border as a rule under the heading
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": "6", "space": "1", "color": "1F3A5F"}.items():
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)


def _para(doc, text, size=9.5, italic=False, color=None):
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(str(text))
    r.font.size, r.italic = Pt(size), italic
    if color:
        r.font.color.rgb = color


def _bullets(doc, items, style="List Bullet", label=None):
    items = [i for i in (items or []) if i]
    if label and items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        r.bold, r.font.size = True, Pt(9.5)
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(it))
        r.font.size = Pt(9.5)


def _page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Page ")
    r.font.size = Pt(8)
    run = p.add_run()
    run.font.size = Pt(8)
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _fmt(v, suffix=""):
    return "–" if v is None else f"{v:,.1f}{suffix}"


def _sec(sections, name):
    s = sections.get(name) or {}
    return {} if "_error" in s else s


# ---------- the report ----------

def build_report(path, facts, sections, val, asking, source_name, model, pages):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)            # A4
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9.5)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("STRICTLY CONFIDENTIAL · INTERNAL FIRST SCREEN · DRAFT")
    hr.font.size, hr.bold, hr.font.color.rgb = Pt(7.5), True, GREY
    _page_number_footer(sec)

    co, d, cur = facts["company"], facts["derived"], facts["currency"]
    name = co.get("name") or "Target company"

    # Title block
    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(0)
    r = t.add_run(f"First Screen: {name}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(17), NAVY
    _para(doc, f"Source: {source_name} ({pages} pp) · Prepared {dt.date.today():%d %b %Y} · "
               f"Processed offline with local model {model} · Figures in {cur} mn unless stated",
          size=8, color=GREY)

    # Verdict banner
    v = _sec(sections, "verdict")
    verdict = v.get("recommendation", "Not assessed")
    b = doc.add_table(rows=1, cols=1)
    b.autofit = False
    c = b.rows[0].cells[0]
    c.width = Cm(17.4)
    _shade(c, VERDICT_FILL.get(verdict, BAND_FILL))
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    rr = p.add_run(f"Initial view: {verdict}.  ")
    rr.bold, rr.font.size = True, Pt(10)
    rr2 = p.add_run(v.get("rationale", ""))
    rr2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Snapshot
    lp = d.get("latest_period")
    cagr = d.get("historical_cagr")
    deal = "; ".join(f"{x.get('item')}: {x.get('value')}" for x in facts["deal_info"][:2]) or "Not stated"
    snap = [
        ["Headquarters", co.get("headquarters"), "Founded", co.get("founded")],
        ["Employees", co.get("employees"), "Business model", co.get("business_model")],
        [f"Revenue ({lp or 'latest'})", _fmt(d.get("latest_revenue")),
         f"{d['ebitda_basis']} margin", _fmt(d["ebitda_margin"].get(lp), "%")],
        ["Revenue CAGR", f"{cagr['pct']}% ({cagr['from']}–{cagr['to']})" if cagr else "–",
         "Deal", deal],
    ]
    st = doc.add_table(rows=0, cols=4)
    st.style = "Table Grid"
    st.autofit = False
    for row in snap:
        cells = st.add_row().cells
        for i, val_ in enumerate(row):
            _cell_text(cells[i], (str(val_)[:160] if val_ else "–"), bold=(i % 2 == 0))
            if i % 2 == 0:
                _shade(cells[i], BAND_FILL)
    _set_widths(st, [3.2, 5.5, 3.2, 5.5])

    # 1. About
    o = _sec(sections, "overview")
    _heading(doc, "1. About the company")
    _para(doc, o.get("about") or co.get("description"))
    _bullets(doc, o.get("investment_highlights"), label="Investment highlights")

    # 2. Revenue mix
    m = _sec(sections, "mix")
    _heading(doc, "2. Revenue mix")
    if facts["revenue_mix"]:
        _table(doc, ["Dimension", "Item", "Value", "Year", "Page"],
               [[x.get("dimension"), x.get("item"), x.get("value"), x.get("year"), x.get("page")]
                for x in facts["revenue_mix"][:8]], [3.2, 6.8, 3, 2.4, 2])
    _para(doc, m.get("revenue_mix_commentary") or "Revenue mix not disclosed in the CIM.")

    # 3. Customer mix
    _heading(doc, "3. Customer mix")
    if facts["customer_mix"]:
        _table(doc, ["Metric", "Value", "Page"],
               [[x.get("metric"), x.get("value"), x.get("page")] for x in facts["customer_mix"][:6]],
               [6.4, 9, 2])
    _para(doc, m.get("customer_mix_commentary") or "Customer data not disclosed in the CIM.")

    # 4. Financial snapshot
    _heading(doc, f"4. Financial snapshot ({cur} mn)")
    periods = _pick_periods(facts["fin_periods"])
    ft = facts["fin_table"]
    if periods:
        from extractor import METRIC_ORDER
        rows, fills = [], []
        for metric in [mt for mt in METRIC_ORDER if mt in ft]:
            rows.append([metric] + [_fmt(ft[metric].get(p, {}).get("value")) for p in periods])
            fills.append(None)
        for label, series in [("Revenue growth", d["growth"]),
                              ("Gross margin", d["gross_margin"]),
                              (f"{d['ebitda_basis']} margin", d["ebitda_margin"])]:
            if series:
                rows.append([label] + [_fmt(series.get(p), "%") for p in periods])
                fills.append([BAND_FILL] * (len(periods) + 1))
        w = 17.4 - 4.2
        _table(doc, ["Metric"] + periods, rows, [4.2] + [w / len(periods)] * len(periods), fills)
        pages_cited = sorted({c["page"] for cells in ft.values() for c in cells.values()
                              if isinstance(c.get("page"), int)})
        _para(doc, f"E = projection. Figures sourced from CIM pp. {', '.join(map(str, pages_cited[:12]))}. "
                   "Growth and margins computed from the extracted figures.", size=7.5, italic=True, color=GREY)
    else:
        _para(doc, "No financial figures could be extracted. Check the CIM's financial section manually.")
    _para(doc, _sec(sections, "financials").get("financial_commentary"))

    # 5. Valuation
    _heading(doc, "5. Comparable valuation")
    if val and val.get("comps"):
        _table(doc, ["Comparable", "EV / Revenue", "EV / EBITDA", "Notes"],
               [[c["company"], _fmt(c["ev_revenue"], "x"), _fmt(c["ev_ebitda"], "x"), c["notes"]]
                for c in val["comps"][:8]], [5.4, 2.8, 2.8, 6.4])
        vrows = []
        for r_ in val["rows"]:
            lo, med, hi = r_["multiples"]
            imp = r_["implied"]
            vrows.append([f"{r_['method']} (n={r_['n']})", f"{lo:.1f}x / {med:.1f}x / {hi:.1f}x",
                          f"{imp[0]:,.0f} – {imp[2]:,.0f} (mid {imp[1]:,.0f})" if imp else "Metric missing"])
        _table(doc, ["Method", "Low / Median / High", f"Implied EV ({cur} mn)"], vrows, [5, 5.2, 7.2])
        _para(doc, f"Applied to {lp or 'latest'} revenue {_fmt(d.get('latest_revenue'))} and "
                   f"{d['ebitda_basis']} {_fmt(d.get('latest_ebitda'))}. Low/high = 25th/75th percentile "
                   "(min/max if fewer than 4 peers). Trading multiples exclude a control premium.",
              size=7.5, italic=True, color=GREY)
    else:
        _para(doc, "No comps file supplied. Add a CSV of peers and multiples from your data terminal "
                   "to value the target (see comps_template.csv).")
        _bullets(doc, v.get("suggested_peers"),
                 label="Model-suggested peer set (unverified - check listing status and relevance)")
    if asking:
        _para(doc, f"Stated price/valuation: {asking['text']} [p. {asking['page']}] → implies "
                   f"{_fmt(asking['ev_revenue'], 'x')} EV/Revenue and {_fmt(asking['ev_ebitda'], 'x')} "
                   f"EV/{d['ebitda_basis']} on {lp} figures (check currency and units).")

    # 6. Synergies
    s = _sec(sections, "synergies")
    _heading(doc, "6. Synergy hypotheses (to test in diligence)")
    _bullets(doc, s.get("revenue_synergies"), label="Revenue")
    _bullets(doc, s.get("cost_synergies"), label="Cost")
    if s.get("strategic_fit"):
        _para(doc, f"Strategic fit: {s['strategic_fit']}")

    # 7. Risks
    rk = _sec(sections, "risks").get("risks") or []
    _heading(doc, "7. Risks and challenges")
    if rk:
        _table(doc, ["Risk", "Severity", "Detail", "Page"],
               [[x.get("title"), x.get("severity"), x.get("detail"), x.get("page")] for x in rk[:6]],
               [4.2, 1.8, 9.8, 1.6],
               fills=[[None, SEVERITY_FILL.get(x.get("severity")), None, None] for x in rk[:6]])
    else:
        _bullets(doc, [f"{x.get('risk')} [p. {x.get('page')}]" for x in facts["risks"][:6]])

    # 8. Diligence questions
    _heading(doc, "8. Key diligence questions")
    _bullets(doc, v.get("diligence_questions"), style="List Number")

    # Gaps & caveats
    gaps = []
    if facts["fin_conflicts"]:
        gaps.append("Conflicting figures in CIM: " + "; ".join(facts["fin_conflicts"][:4]))
    if not facts["customer_mix"]:
        gaps.append("No customer concentration data found.")
    if facts["errors"]:
        gaps.append(f"{len(facts['errors'])} chunk(s) could not be processed; re-run or review manually.")
    gaps.append("Machine-generated first screen from seller-provided material. Verify every figure "
                "against the cited CIM page before circulating.")
    _heading(doc, "Data gaps and caveats")
    for g in gaps:
        _para(doc, g, size=8, italic=True, color=GREY)

    doc.save(path)


def _pick_periods(periods, max_cols=6):
    """Up to 3 latest actual years, LTM, and up to 2 projection years."""
    actual = [p for p in periods if p.startswith("FY") and not p.endswith("E")][-3:]
    ltm = [p for p in periods if p == "LTM"]
    proj = [p for p in periods if p.endswith("E")][:2]
    return (actual + ltm + proj)[:max_cols]
